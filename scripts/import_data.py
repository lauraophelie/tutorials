"""
Script d'import de propriétés immobilières depuis un fichier Excel.

Usage :
    python scripts/import_data.py --config odoo.conf --file scripts/data/properties.xlsx

Ce script :
1. Se connecte à Odoo via odoo.registry (accès direct à l'ORM)
2. Lit le fichier Excel
3. Crée les types, tags et propriétés manquants
4. Logue toutes les opérations dans scripts/logs/
"""

import argparse
import logging
import os
import sys
from datetime import datetime
from pathlib import Path

# Ajoute le dossier Odoo au PYTHONPATH
ODOO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ODOO_ROOT))

import odoo
from odoo.api import Environment
from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# Configuration du logging
# ---------------------------------------------------------------------------

def setup_logging(log_dir: Path) -> logging.Logger:
    """Configure le logger pour écrire dans un fichier et sur la console."""
    log_dir.mkdir(parents=True, exist_ok=True)
    log_file = log_dir / f"import_{datetime.now():%Y-%m-%d_%H-%M-%S}.log"

    logger = logging.getLogger("estate_import")
    logger.setLevel(logging.DEBUG)

    # Handler fichier
    fh = logging.FileHandler(log_file, encoding="utf-8")
    fh.setLevel(logging.DEBUG)
    fh.setFormatter(logging.Formatter(
        "%(asctime)s [%(levelname)s] %(message)s", datefmt="%Y-%m-%d %H:%M:%S"
    ))
    logger.addHandler(fh)

    # Handler console
    ch = logging.StreamHandler()
    ch.setLevel(logging.INFO)
    ch.setFormatter(logging.Formatter("[%(levelname)s] %(message)s"))
    logger.addHandler(ch)

    logger.info(f"Log file: {log_file}")
    return logger


# ---------------------------------------------------------------------------
# Parsing des arguments
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import properties from Excel into Odoo.")
    parser.add_argument("--config", required=True, help="Chemin vers odoo.conf")
    parser.add_argument("--file", required=True, help="Chemin vers le fichier Excel")
    parser.add_argument("--sheet", default="Properties", help="Nom de la feuille Excel")
    parser.add_argument("--dry-run", action="store_true", help="Ne rien écrire en base")

    return parser.parse_args()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def to_bool(value) -> bool:
    """Convertit une valeur Excel en booléen."""
    if value is None:
        return False
    return str(value).strip().lower() in ("oui", "yes", "true", "1", "vrai")


def get_or_create(model, domain, values, logger, dry_run=False):
    """
    Récupère un enregistrement existant ou le crée.
    Retourne (record, created).
    """
    record = model.search(domain, limit=1)
    if record:
        return record, False
    if dry_run:
        logger.info(f"[DRY-RUN] Would create {model._name} with {values}")
        return model.browse(), True
    record = model.create(values)
    logger.info(f"Created {model._name}: {record.display_name} (id={record.id})")
    return record, True


# ---------------------------------------------------------------------------
# Import principal
# ---------------------------------------------------------------------------

def import_properties(env: Environment, excel_path: Path, sheet_name: str, logger: logging.Logger, dry_run: bool = False):
    """Lit le fichier Excel et crée les enregistrements dans Odoo."""

    if not excel_path.exists():
        logger.error(f"Fichier introuvable : {excel_path}")
        return

    wb = load_workbook(excel_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        logger.error(f"Feuille '{sheet_name}' introuvable. Feuilles : {wb.sheetnames}")
        return

    ws = wb[sheet_name]
    headers = [cell.value for cell in ws[1]]
    logger.info(f"Colonnes détectées : {headers}")

    Property = env["estate_property"]
    PropertyType = env["estate.property.type"]
    PropertyTag = env["estate.property.tag"]

    created_count = 0
    skipped_count = 0

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        data = dict(zip(headers, row))

        # Ignore les lignes vides
        if not data.get("name"):
            logger.debug(f"Ligne {row_idx} ignorée (pas de nom)")
            continue

        logger.info(f"--- Ligne {row_idx} : {data['name']} ---")

        try:
            # 1. Résoudre / créer le type
            type_name = data.get("property_type")
            property_type = None
            if type_name:
                property_type, _ = get_or_create(
                    PropertyType,
                    [("name", "=", type_name.strip())],
                    {"name": type_name.strip()},
                    logger, dry_run,
                )

            # 2. Résoudre / créer les tags
            tags = []
            for tag_col in ("tag1", "tag2", "tag3"):
                tag_name = data.get(tag_col)
                if tag_name:
                    tag, _ = get_or_create(
                        PropertyTag,
                        [("name", "=", tag_name.strip())],
                        {"name": tag_name.strip()},
                        logger, dry_run,
                    )
                    tags.append(tag.id)

            # 3. Éviter les doublons sur la propriété
            existing = Property.search([("name", "=", data["name"].strip())], limit=1)
            if existing:
                logger.warning(f"Propriété '{data['name']}' existe déjà (id={existing.id}), ignorée")
                skipped_count += 1
                continue

            # 4. Construire les valeurs de la propriété
            vals = {
                "name": data["name"].strip(),
                "property_type_id": property_type.id if property_type else False,
                "tag_ids": [(6, 0, tags)],
                "postcode": str(data["postcode"]).strip() if data.get("postcode") else False,
                "expected_price": float(data["expected_price"]) if data.get("expected_price") else 0.0,
                "bedrooms": int(data["bedrooms"]) if data.get("bedrooms") else 2,
                "living_area": int(data["living_area"]) if data.get("living_area") else 0,
                "garden": to_bool(data.get("garden")),
                "garden_area": int(data["garden_area"]) if data.get("garden_area") else 0,
                "garden_orientation": data.get("garden_orientation") or False,
            }

            if dry_run:
                logger.info(f"[DRY-RUN] Would create property : {vals}")
            else:
                prop = Property.create(vals)
                logger.info(f"Propriété créée : {prop.name} (id={prop.id})")
                created_count += 1

        except Exception as e:
            logger.error(f"Erreur ligne {row_idx} : {e}", exc_info=True)

    logger.info(f"=== Terminé : {created_count} créées, {skipped_count} ignorées ===")


# ---------------------------------------------------------------------------
# Point d'entrée
# ---------------------------------------------------------------------------

def main():
    args = parse_args()
    logger = setup_logging(Path(__file__).parent / "logs")

    # Charge la config Odoo
    odoo.tools.config.parse_config(["-c", args.config])
    db_name = odoo.tools.config["db_name"]

    if not db_name:
        logger.error("Aucune base de données configurée (db_name). Vérifiez odoo.conf.")
        return

    logger.info(f"Connexion à la base '{db_name}'...")

    # Ouvre un registre Odoo et un environnement
    registry = odoo.registry(db_name)
    with registry.cursor() as cr:
        env = Environment(cr, odoo.SUPERUSER_ID, {})
        try:
            import_properties(env, Path(args.file), args.sheet, logger, dry_run=args.dry_run)
            if not args.dry_run:
                cr.commit()
                logger.info("Transaction commitée.")
            else:
                logger.info("Mode dry-run : rien n'a été écrit.")
        except Exception as e:
            cr.rollback()
            logger.error(f"Erreur fatale, rollback : {e}", exc_info=True)
            raise


if __name__ == "__main__":
    main()