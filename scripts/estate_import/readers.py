"""
Lecteurs de fichiers Excel.

Principe SOLID (S) : une seule responsabilité — lire et parser l'Excel.
Principe SOLID (O) : on peut ajouter un CsvReader sans modifier ExcelReader.
Principe SOLID (D) : les services dépendent de l'interface Reader, pas d'ExcelReader.
"""

from abc import ABC, abstractmethod
from pathlib import Path

from openpyxl import load_workbook

from .models import (
    PropertyDTO, PropertyTypeDTO, PropertyTagDTO, PropertyOfferDTO,
)


class PropertyReader(ABC):
    """
    Interface abstraite pour la lecture de propriétés.

    Toute implémentation (Excel, CSV, JSON...) doit respecter ce contrat.
    """

    @abstractmethod
    def read(self) -> list[PropertyDTO]:
        """Retourne la liste des propriétés lues."""
        raise NotImplementedError


class ExcelPropertyReader(PropertyReader):
    """
    Lit un fichier Excel et retourne des PropertyDTO.

    Le format attendu :
    - Une ligne d'en-tête
    - Une ligne par propriété
    - Les tags dans les colonnes tag1, tag2, tag3
    - Les offres dans les colonnes offer1_partner, offer1_price, etc.
    """

    # Colonnes attendues dans le fichier
    REQUIRED_COLUMNS = {"name"}
    OPTIONAL_COLUMNS = {
        "property_type", "tag1", "tag2", "tag3",
        "offer1_partner", "offer1_price", "offer1_validity",
        "offer2_partner", "offer2_price", "offer2_validity",
        "postcode", "expected_price", "bedrooms", "living_area",
        "garden", "garden_area", "garden_orientation",
    }

    def __init__(self, file_path: Path, sheet_name: str = "Properties"):
        self.file_path = file_path
        self.sheet_name = sheet_name

    def read(self) -> list[PropertyDTO]:
        """Lit le fichier et retourne la liste des propriétés."""
        wb = load_workbook(self.file_path, data_only=True)

        if self.sheet_name not in wb.sheetnames:
            raise ValueError(
                f"Feuille '{self.sheet_name}' introuvable. "
                f"Feuilles disponibles : {wb.sheetnames}"
            )

        ws = wb[self.sheet_name]
        headers = [cell.value for cell in ws[1]]
        self._validate_headers(headers)

        properties = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            data = dict(zip(headers, row))
            if not data.get("name"):
                continue  # ligne vide
            try:
                prop = self._row_to_dto(data, row_idx)
                properties.append(prop)
            except Exception as e:
                raise ValueError(f"Erreur ligne {row_idx} : {e}") from e

        return properties

    def _validate_headers(self, headers: list) -> None:
        """Vérifie que les colonnes obligatoires sont présentes."""
        missing = self.REQUIRED_COLUMNS - set(headers)
        if missing:
            raise ValueError(f"Colonnes obligatoires manquantes : {missing}")

    def _row_to_dto(self, data: dict, row_idx: int) -> PropertyDTO:
        """Convertit une ligne Excel en PropertyDTO."""
        name = str(data["name"]).strip()
        if not name:
            raise ValueError(f"Nom vide ligne {row_idx}")

        # Type
        prop_type = None
        if data.get("property_type"):
            prop_type = PropertyTypeDTO(name=str(data["property_type"]).strip())

        # Tags
        tags = []
        for col in ("tag1", "tag2", "tag3"):
            if data.get(col):
                tags.append(PropertyTagDTO(name=str(data[col]).strip()))

        # Offres
        offers = []
        for i in (1, 2, 3):
            partner = data.get(f"offer{i}_partner")
            price = data.get(f"offer{i}_price")
            if partner and price:
                offers.append(PropertyOfferDTO(
                    partner_name=str(partner).strip(),
                    price=float(price),
                    validity=int(data.get(f"offer{i}_validity") or 7),
                ))

        return PropertyDTO(
            name=name,
            property_type=prop_type,
            tags=tags,
            offers=offers,
            postcode=str(data["postcode"]).strip() if data.get("postcode") else None,
            expected_price=float(data.get("expected_price") or 0.0),
            bedrooms=int(data.get("bedrooms") or 2),
            living_area=int(data.get("living_area") or 0),
            garden=self._to_bool(data.get("garden")),
            garden_area=int(data.get("garden_area") or 0),
            garden_orientation=data.get("garden_orientation") or None,
        )

    @staticmethod
    def _to_bool(value) -> bool:
        """Convertit une valeur Excel en booléen."""
        if value is None:
            return False
        return str(value).strip().lower() in ("oui", "yes", "true", "1", "vrai")